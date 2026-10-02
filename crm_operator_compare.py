"""Compare two operator date-windows from existing daily CRM reports (no LLM)."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import Settings
from crm_batch import build_period_report
from crm_excel_export import (
    CHECKLIST_LABELS,
    HEADER_FILL,
    HEADER_FONT,
    SCORE_LABELS,
    load_daily_reports_in_range,
    _save_workbook_atomic,
)
from models.crm_schemas import CrmAggregateStats, CrmAnalysisReport

SCORE_DELTA_THRESHOLD = 0.15
PCT_DELTA_THRESHOLD = 5.0
QUAL_MIN_B_PCT = 15.0
QUAL_RATIO = 2.0

WRAP = Alignment(wrap_text=True, vertical="top")
PLUS_FILL = PatternFill("solid", fgColor="C6EFCE")
MINUS_FILL = PatternFill("solid", fgColor="FFC7CE")
CTX_FILL = PatternFill("solid", fgColor="DDEBF7")


def _sla_pct(over: int, total: int) -> float | None:
    if total <= 0:
        return None
    return round(max(0.0, min(100.0, 100.0 * (1.0 - over / total))), 1)


def _delta(a: float | None, b: float | None) -> float | None:
    if a is None or b is None:
        return None
    return round(b - a, 2)


def _fmt(value: float | None, *, unit: str) -> str:
    if value is None:
        return "—"
    if unit == "pp":
        return f"{value:.1f}%"
    if unit == "score":
        return f"{value:.2f}"
    if unit == "per_day":
        return f"{value:.1f}"
    return str(value)


def _fmt_delta(value: float | None, *, unit: str) -> str:
    if value is None:
        return "—"
    sign = "+" if value > 0 else ""
    if unit == "pp":
        return f"{sign}{value:.1f} п.п."
    if unit == "score":
        return f"{sign}{value:.2f}"
    if unit == "per_day":
        return f"{sign}{value:.1f}"
    return f"{sign}{value}"


def _slug(name: str) -> str:
    cleaned = re.sub(r"[^\w\-]+", "_", name.strip(), flags=re.UNICODE)
    return cleaned.strip("_") or "op"


def _weekend_days(d_from: date, d_to: date) -> tuple[int, int]:
    total = 0
    weekends = 0
    cur = d_from
    while cur <= d_to:
        total += 1
        if cur.weekday() >= 5:
            weekends += 1
        cur += timedelta(days=1)
    return weekends, total


@dataclass
class MetricDelta:
    key: str
    label: str
    group: str
    a: float | None
    b: float | None
    delta: float | None
    unit: str
    quality: bool


@dataclass
class OperatorCompareResult:
    a_name: str
    b_name: str
    a_from: date
    a_to: date
    b_from: date
    b_to: date
    a_report: CrmAnalysisReport
    b_report: CrmAnalysisReport
    metrics: list[MetricDelta]
    pluses: list[str] = field(default_factory=list)
    minuses: list[str] = field(default_factory=list)
    caveats: list[str] = field(default_factory=list)


def _percent_map(items: list[dict[str, Any]], *, text_key: str) -> dict[str, float]:
    out: dict[str, float] = {}
    for item in items:
        text = str(item.get(text_key) or "").strip()
        if not text:
            continue
        out[text] = float(item.get("percent") or 0.0)
    return out


def classify_quality_findings(
    metrics: list[MetricDelta],
    *,
    a_agg: CrmAggregateStats,
    b_agg: CrmAggregateStats,
    a_name: str,
    b_name: str,
) -> tuple[list[str], list[str]]:
    pluses: list[str] = []
    minuses: list[str] = []
    for m in metrics:
        if not m.quality or m.delta is None:
            continue
        threshold = SCORE_DELTA_THRESHOLD if m.unit == "score" else PCT_DELTA_THRESHOLD
        if m.delta >= threshold:
            pluses.append(f"{m.label}: {_fmt(m.a, unit=m.unit)} → {_fmt(m.b, unit=m.unit)} ({_fmt_delta(m.delta, unit=m.unit)})")
        elif m.delta <= -threshold:
            minuses.append(f"{m.label}: {_fmt(m.a, unit=m.unit)} → {_fmt(m.b, unit=m.unit)} ({_fmt_delta(m.delta, unit=m.unit)})")

    a_err = _percent_map(a_agg.top_errors, text_key="error")
    b_err = _percent_map(b_agg.top_errors, text_key="error")
    a_str = _percent_map(a_agg.top_strengths, text_key="strength")
    b_str = _percent_map(b_agg.top_strengths, text_key="strength")

    for text, b_pct in b_str.items():
        a_pct = a_str.get(text, 0.0)
        if b_pct >= QUAL_MIN_B_PCT and b_pct >= QUAL_RATIO * max(a_pct, 0.01):
            pluses.append(
                f"Сильная сторона «{text}» чаще у {b_name} ({b_pct:.1f}% vs {a_pct:.1f}% у {a_name})"
            )
    for text, b_pct in b_err.items():
        a_pct = a_err.get(text, 0.0)
        if b_pct >= QUAL_MIN_B_PCT and b_pct >= QUAL_RATIO * max(a_pct, 0.01):
            minuses.append(
                f"Ошибка «{text}» чаще у {b_name} ({b_pct:.1f}% vs {a_pct:.1f}% у {a_name})"
            )
    for text, a_pct in a_err.items():
        b_pct = b_err.get(text, 0.0)
        if a_pct >= QUAL_MIN_B_PCT and a_pct >= QUAL_RATIO * max(b_pct, 0.01):
            pluses.append(
                f"Нет типичной ошибки базы «{text}» ({a_pct:.1f}% у {a_name} vs {b_pct:.1f}% у {b_name})"
            )
    return pluses, minuses


def build_compare_metrics(
    a_report: CrmAnalysisReport,
    b_report: CrmAnalysisReport,
) -> list[MetricDelta]:
    a, b = a_report.aggregate, b_report.aggregate
    a_days = a_report.meta.days_in_period or 1
    b_days = b_report.meta.days_in_period or 1
    metrics: list[MetricDelta] = [
        MetricDelta(
            "dialogs_count", "Диалогов (сумма)", "volume",
            float(a_report.meta.dialogs_count), float(b_report.meta.dialogs_count),
            _delta(float(a_report.meta.dialogs_count), float(b_report.meta.dialogs_count)),
            "count", False,
        ),
        MetricDelta(
            "days", "Дней в окне", "volume",
            float(a_days), float(b_days),
            _delta(float(a_days), float(b_days)),
            "count", False,
        ),
        MetricDelta(
            "dialogs_per_day", "Диалогов в день", "volume",
            round(a_report.meta.dialogs_count / a_days, 1),
            round(b_report.meta.dialogs_count / b_days, 1),
            _delta(
                a_report.meta.dialogs_count / a_days,
                b_report.meta.dialogs_count / b_days,
            ),
            "per_day", False,
        ),
        MetricDelta(
            "messages_count", "Сообщений", "volume",
            float(a_report.meta.messages_count), float(b_report.meta.messages_count),
            _delta(float(a_report.meta.messages_count), float(b_report.meta.messages_count)),
            "count", False,
        ),
    ]
    for key, label in SCORE_LABELS.items():
        av, bv = a.avg_scores.get(key), b.avg_scores.get(key)
        metrics.append(MetricDelta(key, label, "score", av, bv, _delta(av, bv), "score", True))
    for key, label in CHECKLIST_LABELS.items():
        av, bv = a.checklist_pass_rate.get(key), b.checklist_pass_rate.get(key)
        metrics.append(MetricDelta(f"cl_{key}", label, "checklist", av, bv, _delta(av, bv), "pp", True))

    a_sla_w = _sla_pct(a.response_time.over_sla_work, a.response_time.responses_count_work)
    b_sla_w = _sla_pct(b.response_time.over_sla_work, b.response_time.responses_count_work)
    a_sla_o = _sla_pct(a.response_time.over_sla_off, a.response_time.responses_count_off)
    b_sla_o = _sla_pct(b.response_time.over_sla_off, b.response_time.responses_count_off)
    metrics += [
        MetricDelta("sla_work", "SLA рабочее (доля ≤2 мин)", "sla", a_sla_w, b_sla_w, _delta(a_sla_w, b_sla_w), "pp", True),
        MetricDelta("sla_off", "SLA вне смены (доля ≤15 мин)", "sla", a_sla_o, b_sla_o, _delta(a_sla_o, b_sla_o), "pp", True),
        MetricDelta(
            "phone_rate", "Телефон → WhatsApp, конверсия", "phone",
            a.phone_success_rate, b.phone_success_rate,
            _delta(a.phone_success_rate, b.phone_success_rate),
            "pp", True,
        ),
        MetricDelta(
            "phone_attempts", "Телефон: попытки", "phone",
            float(a.phone_attempts_total), float(b.phone_attempts_total),
            _delta(float(a.phone_attempts_total), float(b.phone_attempts_total)),
            "count", False,
        ),
        MetricDelta(
            "phone_successes", "Телефон: успехи", "phone",
            float(a.phone_successes_total), float(b.phone_successes_total),
            _delta(float(a.phone_successes_total), float(b.phone_successes_total)),
            "count", False,
        ),
    ]
    return metrics


def _caveats(a_from: date, a_to: date, b_from: date, b_to: date) -> list[str]:
    notes = [
        "Сравниваем качество на диалог, не объём продаж. Диалогов/день — только контекст.",
        f"Пороги: оценка ≥ {SCORE_DELTA_THRESHOLD:.2f}, чеклист/SLA/телефон ≥ {PCT_DELTA_THRESHOLD:.0f} п.п.",
    ]
    a_we, a_n = _weekend_days(a_from, a_to)
    b_we, b_n = _weekend_days(b_from, b_to)
    if b_n > 0 and b_we == b_n and a_n > 0 and a_we < a_n:
        notes.append(
            "Окно замены — выходные, база — смешанные будни и выходные; "
            "микс обращений может отличаться."
        )
    return notes


def compare_operator_windows(
    a_report: CrmAnalysisReport,
    b_report: CrmAnalysisReport,
    *,
    a_name: str,
    b_name: str,
    a_from: date,
    a_to: date,
    b_from: date,
    b_to: date,
) -> OperatorCompareResult:
    metrics = build_compare_metrics(a_report, b_report)
    pluses, minuses = classify_quality_findings(
        metrics,
        a_agg=a_report.aggregate,
        b_agg=b_report.aggregate,
        a_name=a_name,
        b_name=b_name,
    )
    return OperatorCompareResult(
        a_name=a_name,
        b_name=b_name,
        a_from=a_from,
        a_to=a_to,
        b_from=b_from,
        b_to=b_to,
        a_report=a_report,
        b_report=b_report,
        metrics=metrics,
        pluses=pluses,
        minuses=minuses,
        caveats=_caveats(a_from, a_to, b_from, b_to),
    )


def _load_window(settings: Settings, d_from: date, d_to: date) -> CrmAnalysisReport:
    daily = load_daily_reports_in_range(
        settings.output_dir, d_from.isoformat(), d_to.isoformat()
    )
    if not daily:
        raise FileNotFoundError(
            f"Нет дневных crm_report_*.json за {d_from.isoformat()} … {d_to.isoformat()}"
        )
    report = build_period_report(settings, daily, d_from, d_to)
    report.meta.report_type = "operator_window"
    return report


def write_compare_md(result: OperatorCompareResult, path: Path) -> Path:
    a, b = result.a_name, result.b_name
    lines = [
        f"# Сравнение операторов: {a} vs {b}",
        "",
        f"- **{a}:** {result.a_from.isoformat()} — {result.a_to.isoformat()} "
        f"({result.a_report.meta.days_in_period} дн., {result.a_report.meta.dialogs_count} диалогов)",
        f"- **{b}:** {result.b_from.isoformat()} — {result.b_to.isoformat()} "
        f"({result.b_report.meta.days_in_period} дн., {result.b_report.meta.dialogs_count} диалогов)",
        "",
        "## Оговорки",
        "",
    ]
    for note in result.caveats:
        lines.append(f"- {note}")
    lines += [
        "",
        f"## Плюсы {b}",
        "",
    ]
    if result.pluses:
        lines.extend(f"- {item}" for item in result.pluses)
    else:
        lines.append("- Нет метрик выше порога.")
    lines += [
        "",
        f"## Минусы {b}",
        "",
    ]
    if result.minuses:
        lines.extend(f"- {item}" for item in result.minuses)
    else:
        lines.append("- Нет метрик ниже порога.")
    lines += [
        "",
        "## Таблица метрик",
        "",
        f"| Метрика | {a} | {b} | Δ ({b} − {a}) |",
        "|---------|-----|-----|----------------|",
    ]
    for m in result.metrics:
        lines.append(
            f"| {m.label} | {_fmt(m.a, unit=m.unit)} | {_fmt(m.b, unit=m.unit)} | "
            f"{_fmt_delta(m.delta, unit=m.unit)} |"
        )
    lines += ["", "## Топ ошибок", "", f"### {a}", ""]
    for item in result.a_report.aggregate.top_errors[:8]:
        lines.append(f"- ({item['percent']}%) {item['error']}")
    lines += ["", f"### {b}", ""]
    for item in result.b_report.aggregate.top_errors[:8]:
        lines.append(f"- ({item['percent']}%) {item['error']}")
    lines += ["", "## Топ сильных сторон", "", f"### {a}", ""]
    for item in result.a_report.aggregate.top_strengths[:8]:
        lines.append(f"- ({item['percent']}%) {item['strength']}")
    lines += ["", f"### {b}", ""]
    for item in result.b_report.aggregate.top_strengths[:8]:
        lines.append(f"- ({item['percent']}%) {item['strength']}")
    lines.append("")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_compare_excel(result: OperatorCompareResult, path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сравнение"
    a, b = result.a_name, result.b_name
    ws["A1"] = f"Сравнение операторов: {a} vs {b}"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")
    ws["A2"] = (
        f"{a}: {result.a_from.isoformat()} — {result.a_to.isoformat()} | "
        f"{b}: {result.b_from.isoformat()} — {result.b_to.isoformat()}"
    )
    ws.merge_cells("A2:D2")

    headers = ["Метрика", a, b, f"Δ ({b} − {a})"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(4, col, h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for i, m in enumerate(result.metrics, 5):
        ws.cell(i, 1, m.label).alignment = WRAP
        ws.cell(i, 2, _fmt(m.a, unit=m.unit))
        ws.cell(i, 3, _fmt(m.b, unit=m.unit))
        dcell = ws.cell(i, 4, _fmt_delta(m.delta, unit=m.unit))
        if not m.quality:
            for col in range(1, 5):
                ws.cell(i, col).fill = CTX_FILL
        elif m.delta is not None:
            threshold = SCORE_DELTA_THRESHOLD if m.unit == "score" else PCT_DELTA_THRESHOLD
            if m.delta >= threshold:
                dcell.fill = PLUS_FILL
            elif m.delta <= -threshold:
                dcell.fill = MINUS_FILL

    howto = wb.create_sheet("Как читать")
    howto["A1"] = "Как читать сравнение операторов"
    howto["A1"].font = Font(bold=True, size=14)
    howto_lines = [
        "Сравнение двух окон дат: один оператор на календарный день (допущение).",
        "Дневные crm_report_*.json не меняются — это отдельный отчёт поверх готовых дней.",
        "Зелёная дельта — плюс замены (B), красная — минус. Голубые строки — объём, не качество.",
        f"Порог оценки: {SCORE_DELTA_THRESHOLD:.2f} балла. Порог чеклиста/SLA/телефона: {PCT_DELTA_THRESHOLD:.0f} п.п.",
        "Качественные плюсы/минусы: паттерн ≥15% диалогов замены и ≥2× доли базы.",
        "Нельзя читать «кто продал больше»: разный объём и микс дней (будни vs выходные).",
    ]
    howto_lines.extend(result.caveats)
    for i, line in enumerate(howto_lines, 3):
        howto[f"A{i}"] = line
        howto[f"A{i}"].alignment = WRAP
    howto.column_dimensions["A"].width = 110

    plus_sheet = wb.create_sheet("Плюсы и минусы")
    plus_sheet["A1"] = f"Плюсы {b}"
    plus_sheet["A1"].font = Font(bold=True, size=12)
    plus_sheet["A1"].fill = PLUS_FILL
    row = 2
    for item in result.pluses or ["Нет метрик выше порога."]:
        plus_sheet[f"A{row}"] = item
        plus_sheet[f"A{row}"].alignment = WRAP
        row += 1
    row += 1
    plus_sheet[f"A{row}"] = f"Минусы {b}"
    plus_sheet[f"A{row}"].font = Font(bold=True, size=12)
    plus_sheet[f"A{row}"].fill = MINUS_FILL
    row += 1
    for item in result.minuses or ["Нет метрик ниже порога."]:
        plus_sheet[f"A{row}"] = item
        plus_sheet[f"A{row}"].alignment = WRAP
        row += 1
    plus_sheet.column_dimensions["A"].width = 110

    for col, width in enumerate([42, 22, 22, 22], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A5"
    return _save_workbook_atomic(wb, path)


def compare_stem(a_name: str, a_from: date, a_to: date, b_name: str, b_from: date, b_to: date) -> str:
    return (
        f"CRM_OPERATOR_COMPARE_{_slug(a_name)}_{a_from.isoformat()}_{a_to.isoformat()}"
        f"_vs_{_slug(b_name)}_{b_from.isoformat()}_{b_to.isoformat()}"
    )


def run_operator_compare(
    settings: Settings,
    *,
    a_name: str,
    a_from: date,
    a_to: date,
    b_name: str,
    b_from: date,
    b_to: date,
) -> dict[str, str]:
    if a_from > a_to or b_from > b_to:
        raise ValueError("date_from must be <= date_to for each operator window")
    a_report = _load_window(settings, a_from, a_to)
    b_report = _load_window(settings, b_from, b_to)
    result = compare_operator_windows(
        a_report,
        b_report,
        a_name=a_name,
        b_name=b_name,
        a_from=a_from,
        a_to=a_to,
        b_from=b_from,
        b_to=b_to,
    )
    stem = compare_stem(a_name, a_from, a_to, b_name, b_from, b_to)
    md_path = settings.project_root / "docs" / "analysis" / f"{stem}.md"
    xlsx_path = settings.output_dir / "crm_excel" / f"{stem}.xlsx"
    write_compare_md(result, md_path)
    write_compare_excel(result, xlsx_path)
    return {"md": str(md_path), "excel": str(xlsx_path)}
